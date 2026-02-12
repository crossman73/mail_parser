"""Word(DOCX) 내보내기"""

from datetime import datetime
from io import BytesIO
from typing import Any, Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from .base import TimelineExporter


class DocxExporter(TimelineExporter):
    """타임라인을 DOCX로 내보내기"""

    @property
    def format_name(self) -> str:
        return 'Word'

    @property
    def content_type(self) -> str:
        return 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'

    @property
    def file_extension(self) -> str:
        return 'docx'

    def export(self, timeline: Dict[str, Any]) -> BytesIO:
        doc = Document()

        # 스타일 설정
        style = doc.styles['Normal']
        style.font.name = 'Malgun Gothic'
        style.font.size = Pt(11)

        # 제목
        title = doc.add_heading(timeline.get('title', '타임라인'), level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 부제목/설명
        desc = timeline.get('description', '')
        if desc:
            p = doc.add_paragraph(desc)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph('')  # 빈 줄

        # 요약 정보
        doc.add_heading('1. 요약', level=1)
        events = timeline.get('events', [])
        key_count = sum(1 for e in events if e.get('is_key_event'))

        summary_table = doc.add_table(rows=4 + (1 if key_count else 0), cols=2)
        summary_table.style = 'Table Grid'
        summary_table.columns[0].width = Inches(2.5)
        summary_table.columns[1].width = Inches(4.0)

        rows = [
            ('총 이벤트', f'{len(events)}건'),
            ('기간', f"{timeline.get('date_range_start', 'N/A')} ~ {timeline.get('date_range_end', 'N/A')}"),
            ('상태', timeline.get('status', 'draft')),
            ('생성일', datetime.now().strftime('%Y-%m-%d')),
        ]
        if key_count:
            rows.append(('핵심 이벤트', f'{key_count}건'))

        for i, (label, value) in enumerate(rows):
            cells = summary_table.rows[i].cells
            cells[0].text = label
            cells[1].text = value
            # 라벨 볼드
            for run in cells[0].paragraphs[0].runs:
                run.bold = True

        doc.add_paragraph('')

        # 이벤트 상세
        doc.add_heading('2. 이벤트 상세', level=1)

        for idx, evt in enumerate(events, 1):
            ts = evt.get('timestamp', '')
            if isinstance(ts, str) and len(ts) > 10:
                ts = ts[:10]
            marker = '★ ' if evt.get('is_key_event') else ''

            # 이벤트 제목
            h = doc.add_heading(f"{marker}{idx}. [{ts}] {evt.get('title', '')}", level=2)

            # 유형
            etype = evt.get('event_type', 'email')
            p_type = doc.add_paragraph()
            run = p_type.add_run(f'유형: ')
            run.bold = True
            p_type.add_run(etype)

            # 설명
            desc = evt.get('description', '')
            if desc:
                p = doc.add_paragraph()
                run = p.add_run('설명: ')
                run.bold = True
                p.add_run(desc)

            # 참여자
            participants = evt.get('participants', [])
            if participants:
                p = doc.add_paragraph()
                run = p.add_run('참여자: ')
                run.bold = True
                p.add_run(', '.join(participants))

            # 법적 중요성
            legal = evt.get('legal_significance', '')
            if legal:
                p = doc.add_paragraph()
                run = p.add_run('법적 중요성: ')
                run.bold = True
                run_content = p.add_run(legal)
                run_content.font.color.rgb = RGBColor(0x00, 0x66, 0x99)

            # 비고
            notes = evt.get('notes', '')
            if notes:
                p = doc.add_paragraph()
                run = p.add_run('비고: ')
                run.bold = True
                p.add_run(notes)

            # 첨부파일
            attachments = evt.get('attachments', [])
            if attachments:
                p = doc.add_paragraph()
                run = p.add_run('첨부파일: ')
                run.bold = True
                p.add_run(', '.join(attachments))

            doc.add_paragraph('')  # 이벤트 간 간격

        # 꼬리말
        doc.add_paragraph('')
        footer = doc.add_paragraph(
            f"이 문서는 {datetime.now().strftime('%Y-%m-%d %H:%M')}에 자동 생성되었습니다.")
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf
