"""PowerPoint(PPTX) 내보내기"""

from datetime import datetime
from io import BytesIO
from typing import Any, Dict

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .base import TimelineExporter


class PptxExporter(TimelineExporter):
    """타임라인을 PPTX로 내보내기"""

    @property
    def format_name(self) -> str:
        return 'PowerPoint'

    @property
    def content_type(self) -> str:
        return 'application/vnd.openxmlformats-officedocument.presentationml.presentation'

    @property
    def file_extension(self) -> str:
        return 'pptx'

    def export(self, timeline: Dict[str, Any]) -> BytesIO:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        # 표지 슬라이드
        self._add_title_slide(prs, timeline)

        # 요약 슬라이드
        self._add_summary_slide(prs, timeline)

        # 이벤트 슬라이드 (5개씩 그룹)
        events = timeline.get('events', [])
        for i in range(0, len(events), 5):
            chunk = events[i:i + 5]
            self._add_events_slide(prs, chunk, i + 1, len(events))

        # 종합 슬라이드
        self._add_closing_slide(prs, timeline)

        buf = BytesIO()
        prs.save(buf)
        buf.seek(0)
        return buf

    def _add_title_slide(self, prs, timeline):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 레이아웃

        # 배경색 (진한 남색)
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

        # 제목
        txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = timeline.get('title', '타임라인')
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        p.alignment = PP_ALIGN.CENTER

        # 부제목
        sub = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(11), Inches(1))
        tf2 = sub.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        desc = timeline.get('description', '')
        p2.text = desc if desc else '법원 증거 타임라인'
        p2.font.size = Pt(20)
        p2.font.color.rgb = RGBColor(0xbb, 0xbb, 0xbb)
        p2.alignment = PP_ALIGN.CENTER

        # 날짜
        date_box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(11), Inches(0.8))
        tf3 = date_box.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = f"생성일: {datetime.now().strftime('%Y-%m-%d')}"
        p3.font.size = Pt(14)
        p3.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        p3.alignment = PP_ALIGN.CENTER

    def _add_summary_slide(self, prs, timeline):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # 제목
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = '타임라인 요약'
        p.font.size = Pt(28)
        p.font.bold = True

        events = timeline.get('events', [])
        summary_items = [
            f"총 이벤트 수: {len(events)}건",
            f"기간: {timeline.get('date_range_start', 'N/A')} ~ {timeline.get('date_range_end', 'N/A')}",
            f"상태: {timeline.get('status', 'draft')}",
        ]

        # 핵심 이벤트 수
        key_count = sum(1 for e in events if e.get('is_key_event'))
        if key_count:
            summary_items.append(f"핵심 이벤트: {key_count}건")

        content_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
        tf2 = content_box.text_frame
        tf2.word_wrap = True
        for item in summary_items:
            p = tf2.add_paragraph() if tf2.paragraphs[0].text else tf2.paragraphs[0]
            p.text = f"• {item}"
            p.font.size = Pt(18)
            p.space_after = Pt(12)

    def _add_events_slide(self, prs, events, start_idx, total):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # 제목
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.6))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        end_idx = min(start_idx + len(events) - 1, total)
        p.text = f'이벤트 {start_idx}-{end_idx} / {total}'
        p.font.size = Pt(20)
        p.font.bold = True

        y_offset = 1.2
        for evt in events:
            # 이벤트 카드
            box = slide.shapes.add_textbox(
                Inches(0.5), Inches(y_offset), Inches(12), Inches(1.0))
            tf2 = box.text_frame
            tf2.word_wrap = True

            # 날짜 + 제목
            p_title = tf2.paragraphs[0]
            ts = evt.get('timestamp', '')
            if isinstance(ts, str) and len(ts) > 10:
                ts = ts[:10]
            marker = '★ ' if evt.get('is_key_event') else ''
            p_title.text = f"{marker}[{ts}] {evt.get('title', '')}"
            p_title.font.size = Pt(16)
            p_title.font.bold = True

            # 설명
            desc = evt.get('description', '')
            if desc:
                p_desc = tf2.add_paragraph()
                p_desc.text = desc[:200]
                p_desc.font.size = Pt(12)
                p_desc.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            # 법적 중요성
            legal = evt.get('legal_significance', '')
            if legal:
                p_legal = tf2.add_paragraph()
                p_legal.text = f"⚖ {legal}"
                p_legal.font.size = Pt(11)
                p_legal.font.color.rgb = RGBColor(0x00, 0x66, 0x99)

            y_offset += 1.2

    def _add_closing_slide(self, prs, timeline):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

        box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = '— 끝 —'
        p.font.size = Pt(36)
        p.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = f"이 문서는 {datetime.now().strftime('%Y-%m-%d %H:%M')}에 자동 생성되었습니다."
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        p2.alignment = PP_ALIGN.CENTER
